import io
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

from resume_parser.base import Resume


ACCENT = RGBColor(0x24, 0x5C, 0xA6)  # blue accent


def _tight(p):
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1
    try:
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    except Exception:
        # Fallback if environment lacks enum value support
        pass
    return p


def _heading(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = ACCENT
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _tight(p)


def _muted(doc: Document, text: str):
    p = doc.add_paragraph()
    for r in p.runs:
        r.font.size = Pt(9)
    return _tight(p)


def _small(doc: Document, text: str):
    p = doc.add_paragraph(text)
    for r in p.runs:
        r.font.size = Pt(9)
    return _tight(p)


def _lines(doc: Document, lines):
    for b in lines:
        if not b:
            continue
        para = doc.add_paragraph(b)
        for r in para.runs:
            r.font.size = Pt(10)
        _tight(para)


def generate_modern_docx(resume: Resume) -> io.BytesIO:
    doc = Document()

    # Header
    p = doc.add_paragraph()
    r = p.add_run(resume.name or "Candidate")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = ACCENT
    _tight(p)

    if resume.headline:
        h = doc.add_paragraph(resume.headline)
        for rr in h.runs:
            rr.font.size = Pt(11)
        _tight(h)

    contact_parts = []
    c = resume.contact
    for part in [c.email, c.phone, c.location, c.linkedin, c.github, c.website]:
        if part:
            contact_parts.append(part)
    if contact_parts:
        cc = doc.add_paragraph(" · ".join(contact_parts))
        for rr in cc.runs:
            rr.font.size = Pt(9)
        _tight(cc)

    # Summary
    if resume.summary:
        _heading(doc, "Summary")
        s = doc.add_paragraph(resume.summary)
        for rr in s.runs:
            rr.font.size = Pt(10)
        _tight(s)

    # Skills
    if resume.skills:
        _heading(doc, "Skills")
        s = doc.add_paragraph(", ".join(resume.skills))
        for rr in s.runs:
            rr.font.size = Pt(10)
        _tight(s)

    # Experience
    if resume.experience:
        _heading(doc, "Experience")
        for e in resume.experience:
            header = e.job_title
            if e.company:
                header += f" at {e.company}"
            if e.start_date or e.end_date:
                header += f" [{(e.start_date or '').strip()} - {(e.end_date or '').strip()}]"
            h = doc.add_paragraph(header)
            for rr in h.runs:
                rr.font.size = Pt(11)
                # no bold for entries per request
            _tight(h)
            _lines(doc, e.bullets)

    # Education
    if resume.education:
        _heading(doc, "Education")
        for edu in resume.education:
            header = edu.institution
            if edu.degree:
                header += f" — {edu.degree}"
            if edu.start_date or edu.end_date:
                header += f" [{(edu.start_date or '').strip()} - {(edu.end_date or '').strip()}]"
            h = doc.add_paragraph(header)
            for rr in h.runs:
                rr.font.size = Pt(11)
                # no bold for entries per request
            _tight(h)
            _lines(doc, edu.details)

    # Projects
    if resume.projects:
        _heading(doc, "Projects")
        for p in resume.projects:
            title = p.name
            if p.link:
                title += f" ({p.link})"
            tr = doc.add_paragraph(title)
            for rr in tr.runs:
                rr.font.size = Pt(11)
                # no bold for entries per request
            _tight(tr)
            _lines(doc, p.bullets)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
