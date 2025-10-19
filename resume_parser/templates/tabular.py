import io
from typing import List
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from resume_parser.base import Resume, ExperienceEntry, EducationEntry, ProjectEntry


def _add_cell_text(cell, text: str, bold: bool = False, size: int = 10):
    p = cell.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def _bullets(cell, bullets: List[str]):
    for b in bullets:
        if not b:
            continue
        para = cell.add_paragraph(b, style="List Bullet")
        for r in para.runs:
            r.font.size = Pt(10)


def generate_tabular_docx(resume: Resume) -> io.BytesIO:
    doc = Document()

    # Name header
    p = doc.add_paragraph()
    r = p.add_run(resume.name or "Candidate")
    r.bold = True
    r.font.size = Pt(22)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if resume.headline:
        hp = doc.add_paragraph(resume.headline)
        for rr in hp.runs:
            rr.font.size = Pt(11)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Contact row table
    contact_tbl = doc.add_table(rows=1, cols=1)
    contact_tbl.autofit = True
    contact_texts = []
    c = resume.contact
    for part in [c.email, c.phone, c.location, c.linkedin, c.github, c.website]:
        if part:
            contact_texts.append(part)
    if contact_texts:
        _add_cell_text(contact_tbl.rows[0].cells[0], " | ".join(contact_texts), size=9)

    # Skills row
    if resume.skills:
        skills_tbl = doc.add_table(rows=1, cols=1)
        skills_tbl.autofit = True
        _add_cell_text(skills_tbl.rows[0].cells[0], "Skills", bold=True, size=12)
        _add_cell_text(skills_tbl.rows[0].cells[0], ", ".join(resume.skills), size=10)

    # Experience table
    if resume.experience:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        _add_cell_text(hdr[0], "Title", bold=True)
        _add_cell_text(hdr[1], "Company / Dates", bold=True)
        _add_cell_text(hdr[2], "Highlights", bold=True)
        for e in resume.experience:
            row = t.add_row().cells
            _add_cell_text(row[0], e.job_title or "Experience", bold=True)
            company_dates = []
            if e.company:
                company_dates.append(e.company)
            if e.start_date or e.end_date:
                company_dates.append(f"{(e.start_date or '').strip()} - {(e.end_date or '').strip()}")
            _add_cell_text(row[1], " | ".join([p for p in company_dates if p]))
            _bullets(row[2], e.bullets)

    # Education table
    if resume.education:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        _add_cell_text(hdr[0], "Institution", bold=True)
        _add_cell_text(hdr[1], "Degree / Dates", bold=True)
        _add_cell_text(hdr[2], "Details", bold=True)
        for edu in resume.education:
            row = t.add_row().cells
            _add_cell_text(row[0], edu.institution, bold=True)
            deg_dates = []
            if edu.degree:
                deg_dates.append(edu.degree)
            if edu.start_date or edu.end_date:
                deg_dates.append(f"{(edu.start_date or '').strip()} - {(edu.end_date or '').strip()}")
            _add_cell_text(row[1], " | ".join([p for p in deg_dates if p]))
            _bullets(row[2], edu.details)

    # Projects table
    if resume.projects:
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        _add_cell_text(hdr[0], "Project", bold=True)
        _add_cell_text(hdr[1], "Highlights", bold=True)
        for p in resume.projects:
            row = t.add_row().cells
            title = p.name
            if p.link:
                title += f" ({p.link})"
            _add_cell_text(row[0], title, bold=True)
            _bullets(row[1], p.bullets)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
