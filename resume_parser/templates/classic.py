import io
from typing import Optional
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches, RGBColor

from resume_parser.base import Resume, ExperienceEntry, EducationEntry, ProjectEntry


def _add_heading(doc: Document, text: str, size: int = 16):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _add_separator(doc: Document):
    doc.add_paragraph("_" * 60)


def _add_key_value_line(doc: Document, label: str, value: Optional[str]):
    if not value:
        return
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)


def _add_bullets(doc: Document, bullets):
    for b in bullets:
        if not b:
            continue
        doc.add_paragraph(b, style="List Bullet")


def generate_classic_docx(resume: Resume) -> io.BytesIO:
    doc = Document()

    # Name and headline
    name_p = doc.add_paragraph()
    name_run = name_p.add_run(resume.name or "Candidate")
    name_run.bold = True
    name_run.font.size = Pt(22)

    if resume.headline:
        h = doc.add_paragraph(resume.headline)
        h.runs[0].font.size = Pt(12)

    # Contact
    c = resume.contact
    contact_line_parts = []
    for part in [c.email, c.phone, c.location, c.linkedin, c.github, c.website]:
        if part:
            contact_line_parts.append(part)
    if contact_line_parts:
        cl = doc.add_paragraph(" | ".join(contact_line_parts))
        for r in cl.runs:
            r.font.size = Pt(10)

    _add_separator(doc)

    # Summary
    if resume.summary:
        _add_heading(doc, "Summary", 14)
        doc.add_paragraph(resume.summary)

    # Skills
    if resume.skills:
        _add_heading(doc, "Skills", 14)
        doc.add_paragraph(", ".join(resume.skills))

    # Experience
    if resume.experience:
        _add_heading(doc, "Experience", 14)
        for e in resume.experience:
            header = e.job_title
            if e.company:
                header += f" — {e.company}"
            dates = ""
            if e.start_date or e.end_date:
                dates = f" ({(e.start_date or '').strip()} - {(e.end_date or '').strip()})"
            h = doc.add_paragraph()
            hr = h.add_run(header + dates)
            hr.bold = True
            _add_bullets(doc, e.bullets)

    # Education
    if resume.education:
        _add_heading(doc, "Education", 14)
        for edu in resume.education:
            header = edu.institution
            if edu.degree:
                header += f" — {edu.degree}"
            dates = ""
            if edu.start_date or edu.end_date:
                dates = f" ({(edu.start_date or '').strip()} - {(edu.end_date or '').strip()})"
            h = doc.add_paragraph()
            hr = h.add_run(header + dates)
            hr.bold = True
            _add_bullets(doc, edu.details)

    # Projects
    if resume.projects:
        _add_heading(doc, "Projects", 14)
        for p in resume.projects:
            title = p.name
            if p.link:
                title += f" ({p.link})"
            tr = doc.add_paragraph().add_run(title)
            tr.bold = True
            _add_bullets(doc, p.bullets)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
