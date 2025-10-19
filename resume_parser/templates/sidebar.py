import io
from typing import List
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from resume_parser.base import Resume


def _set_run_style(run, size=10, bold=False):
    run.font.size = Pt(size)
    run.bold = bold


def _add_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    _set_run_style(r, size=11, bold=True)
    return p


def _add_bullets(container, bullets: List[str]):
    for b in bullets:
        if not b:
            continue
        para = container.add_paragraph(b, style="List Bullet")
        for r in para.runs:
            r.font.size = Pt(10)


def generate_sidebar_docx(resume: Resume) -> io.BytesIO:
    doc = Document()

    # Create a two-column layout using a 1-row, 2-column table
    layout = doc.add_table(rows=1, cols=2)
    layout.autofit = True
    left, right = layout.rows[0].cells

    # LEFT: Sidebar - contact and skills
    name_p = left.add_paragraph()
    name_run = name_p.add_run(resume.name or "Candidate")
    _set_run_style(name_run, size=18, bold=True)

    if resume.headline:
        hl = left.add_paragraph(resume.headline)
        for r in hl.runs:
            r.font.size = Pt(10)

    # Contact
    c = resume.contact
    contact_parts = []
    for part in [c.email, c.phone, c.location, c.linkedin, c.github, c.website]:
        if part:
            contact_parts.append(part)
    if contact_parts:
        cp = left.add_paragraph("\n".join(contact_parts))
        for r in cp.runs:
            r.font.size = Pt(9)

    # Skills
    if resume.skills:
        _add_heading(left, "Skills")
        sp = left.add_paragraph("\n".join(resume.skills))
        for r in sp.runs:
            r.font.size = Pt(9)

    # RIGHT: Main content - Summary, Experience, Education, Projects
    if resume.summary:
        _add_heading(right, "Summary")
        s = right.add_paragraph(resume.summary)
        for r in s.runs:
            r.font.size = Pt(10)

    if resume.experience:
        _add_heading(right, "Experience")
        for e in resume.experience:
            header = e.job_title
            if e.company:
                header += f" — {e.company}"
            if e.start_date or e.end_date:
                header += f" [{(e.start_date or '').strip()} - {(e.end_date or '').strip()}]"
            h = right.add_paragraph(header)
            for rr in h.runs:
                rr.font.size = Pt(11)
                rr.bold = True
            _add_bullets(right, e.bullets)

    if resume.education:
        _add_heading(right, "Education")
        for edu in resume.education:
            header = edu.institution
            if edu.degree:
                header += f" — {edu.degree}"
            if edu.start_date or edu.end_date:
                header += f" [{(edu.start_date or '').strip()} - {(edu.end_date or '').strip()}]"
            h = right.add_paragraph(header)
            for rr in h.runs:
                rr.font.size = Pt(11)
                rr.bold = True
            _add_bullets(right, edu.details)

    if resume.projects:
        _add_heading(right, "Projects")
        for p in resume.projects:
            title = p.name
            if p.link:
                title += f" ({p.link})"
            tr = right.add_paragraph(title)
            for rr in tr.runs:
                rr.font.size = Pt(11)
                rr.bold = True
            _add_bullets(right, p.bullets)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
