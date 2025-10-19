import io
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from resume_parser.base import Resume


def _hr(doc: Document):
    doc.add_paragraph("—" * 50)


def _heading(doc: Document, text: str):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    return p


def _bullets(doc: Document, items):
    for it in items:
        if not it:
            continue
        para = doc.add_paragraph(f"• {it}")
        for r in para.runs:
            r.font.size = Pt(10)


def generate_minimal_docx(resume: Resume) -> io.BytesIO:
    doc = Document()

    # Centered name and headline
    p = doc.add_paragraph()
    r = p.add_run(resume.name or "Candidate")
    r.bold = True
    r.font.size = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if resume.headline:
        hp = doc.add_paragraph(resume.headline)
        for rr in hp.runs:
            rr.font.size = Pt(10)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Contact single line
    c = resume.contact
    parts = [x for x in [c.email, c.phone, c.location, c.linkedin, c.github, c.website] if x]
    if parts:
        cp = doc.add_paragraph(" • ".join(parts))
        for rr in cp.runs:
            rr.font.size = Pt(9)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _hr(doc)

    if resume.summary:
        _heading(doc, "Summary")
        s = doc.add_paragraph(resume.summary)
        for rr in s.runs:
            rr.font.size = Pt(10)

    if resume.skills:
        _heading(doc, "Skills")
        s = doc.add_paragraph(", ".join(resume.skills))
        for rr in s.runs:
            rr.font.size = Pt(10)

    if resume.experience:
        _heading(doc, "Experience")
        for e in resume.experience:
            header = e.job_title
            if e.company:
                header += f" — {e.company}"
            if e.start_date or e.end_date:
                header += f" [{(e.start_date or '').strip()} - {(e.end_date or '').strip()}]"
            h = doc.add_paragraph(header)
            for rr in h.runs:
                rr.font.size = Pt(11)
                rr.bold = True
            _bullets(doc, e.bullets)

    if resume.education:
        _heading(doc, "Education")
        for edu in resume.education:
            header = edu.institution
            if edu.degree:
                header += f" — {edu.degree}"
            if edu.start_date or edu.end_date:
                header += f" [{(edu.start_date or '').strip()} - {(edu.end_date or '').strip()}]"
            h = doc.add_paragraph(header)
            for rr in h.runs:
                rr.font.size = Pt(11)
                rr.bold = True
            _bullets(doc, edu.details)

    if resume.projects:
        _heading(doc, "Projects")
        for p in resume.projects:
            title = p.name
            if p.link:
                title += f" ({p.link})"
            tr = doc.add_paragraph(title)
            for rr in tr.runs:
                rr.font.size = Pt(11)
                rr.bold = True
            _bullets(doc, p.bullets)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
