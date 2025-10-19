import io
from typing import IO, List
from pdfminer.high_level import extract_text as pdf_extract_text
from docx import Document as DocxDocument


def _extract_pdf_text(file_stream: IO[bytes]) -> str:
    """Extract text from PDF, attempting to preserve multi-column separation.

    Strategy: use pdfminer text extraction with LAParams tuned for columns, then
    split lines and heuristically rejoin by detecting clear column separators.
    If we detect two columns (wide line with a long run of spaces), we split
    each line at the separator and append left then right column lines
    separately to avoid interleaving content.
    """
    from pdfminer.layout import LAParams

    file_stream.seek(0)
    try:
        laparams = LAParams()
        laparams.line_margin = 0.1  # tighter line grouping
        laparams.char_margin = 2.0  # allow some spacing within words
        laparams.word_margin = 0.1
        laparams.detect_vertical = False
        # extract with parameters
        text = pdf_extract_text(file_stream, laparams=laparams) or ""
    except Exception:
        file_stream.seek(0)
        text = pdf_extract_text(file_stream) or ""

    if not text:
        return ""

    lines = text.splitlines()
    # Heuristic: detect a consistent column split index by scanning for lines that
    # contain long spaces separating two non-empty parts
    sep_counts = {}
    for line in lines:
        # find a long stretch of spaces as a potential column gap
        if '     ' in line:  # five or more spaces
            # find the widest gap index
            longest = 0
            longest_idx = None
            current_len = 0
            for i, ch in enumerate(line):
                if ch == ' ':
                    current_len += 1
                    if current_len > longest:
                        longest = current_len
                        longest_idx = i - current_len + 1
                else:
                    current_len = 0
            if longest >= 5 and longest_idx is not None:
                sep_counts[longest_idx] = sep_counts.get(longest_idx, 0) + 1

    if not sep_counts:
        return text

    # choose the most frequent split position
    split_idx = max(sep_counts.items(), key=lambda kv: kv[1])[0]

    left_col: List[str] = []
    right_col: List[str] = []
    for raw in lines:
        if len(raw) > split_idx + 3 and raw[split_idx:split_idx + 3].isspace():
            left = raw[:split_idx].rstrip()
            right = raw[split_idx:].strip()
            if left:
                left_col.append(left)
            if right:
                right_col.append(right)
        else:
            # Unknown alignment: keep line break between columns to avoid mixing
            left_col.append(raw)
            left_col.append("")

    # Return left column, blank line, then right column. This prevents
    # interleaving certifications with experience from adjacent column.
    normalized = "\n".join([*left_col, "", *right_col])
    return normalized


def _extract_docx_text(file_stream: IO[bytes]) -> str:
    file_stream.seek(0)
    doc = DocxDocument(file_stream)
    parts = []
    for p in doc.paragraphs:
        if p.text:
            parts.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    parts.append(cell_text)
    return "\n".join(parts)


def extract_text_from_upload(file_stream: IO[bytes], filename: str) -> str:
    name = filename.lower().strip()
    if name.endswith(".pdf"):
        return _extract_pdf_text(file_stream)
    if name.endswith(".docx"):
        return _extract_docx_text(file_stream)
    raise ValueError("Unsupported file type; only PDF and DOCX are supported.")
